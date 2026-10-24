#!/usr/bin/env python3
"""
Prognos · Black and White — ingestão do arquivo técnico.

Lê uma pasta (ex.: G:\\1 - Arquivo Tecnico), extrai o texto de cada arquivo,
indexa na base da planta (Supabase) e, se você pedir, usa a IA para extrair
eventos de manutenção (falhas, trocas, medições) de relatórios e históricos.

Os arquivos originais nunca saem do seu computador: só o texto extraído vai
para a base. O programa pode ser interrompido (Ctrl+C) e retomado: ele lembra
o que já leu.

Uso típico (Windows):  iniciar.bat
Uso manual:            python prognos_ingest.py "G:\\1 - Arquivo Tecnico"
Só ver o que existe:   python prognos_ingest.py "G:\\..." --inventario
Teste sem enviar nada: python prognos_ingest.py "G:\\..." --seco --limite 50
Com IA (com teto):     python prognos_ingest.py "G:\\..." --ia --orcamento-usd 10
"""
import argparse, datetime as dt, getpass, io, json, os, re, sqlite3, sys, time, zipfile
from pathlib import Path

SB_URL = "https://hxzgorearzxtvippjhjg.supabase.co"
SB_KEY = "sb_publishable_eIfvP_NdNxo8nXYxookWiw_MTmQkCxY"
HOME = Path.home() / ".prognos"
HOME.mkdir(exist_ok=True)

TEXTO = {".txt", ".csv", ".log", ".md", ".json", ".xml", ".htm", ".html", ".st", ".scl", ".awl", ".stl", ".l5x", ".l5k", ".il", ".exp", ".iec", ".prg", ".udt", ".ini", ".cfg"}
IGNORAR = {".exe", ".dll", ".msi", ".bak", ".tmp", ".lnk", ".sys", ".bin", ".iso", ".dat", ".db", ".ctb", ".dwt", ".dwl", ".dwl2", ".sv$", ".ds_store", ".ini~", ".pyc"}
IMAGEM = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
DESENHO = {".dwg", ".dxf", ".dgn", ".step", ".stp", ".iges", ".igs", ".sldprt", ".sldasm", ".ipt", ".iam", ".prt", ".asm"}
VIDEO = {".mp4", ".avi", ".mov", ".wmv", ".mkv", ".3gp"}
COMPACTADO = {".zip", ".7z", ".rar"}
DISCIPLINA = {"CIV": "civil", "MEC": "mecânica", "ELE": "elétrica", "PCD": "processo", "TUB": "tubulação", "CTO": "controle/automação", "ARQ": "arquitetura", "INS": "instrumentação", "EST": "estrutura", "HID": "hidráulica"}
HIST_RE = re.compile(r"\b(falh|quebr|trinc|troc|substitu|danific|travad|rompeu|queim|desarm|vazament|superaquec|vibra[cç]|ordem de manuten|nota de manuten|\bOM\b|relat[oó]rio de (turno|ocorr|falha|inspe)|an[aá]lise de falha|causa raiz|RCA|parada n[aã]o programada|corretiva|preventiva|preditiva)", re.I)
MANUAL_PASTAS = re.compile(r"apostila|manual|catalog|cat[aá]logo|norma|padr[oõ]es|tabelas|blocos", re.I)


def log(*a):
    print(*a, flush=True)


# ---------------------------------------------------------------- estado local
class Estado:
    def __init__(self, pasta):
        nome = re.sub(r"[^\w]+", "_", str(pasta))[-60:]
        self.db = sqlite3.connect(HOME / f"estado_{nome}.sqlite")
        self.db.execute("create table if not exists arq(caminho text primary key, tamanho int, mtime real, status text, rid int, info text)")
        self.db.execute("create table if not exists gasto(id integer primary key, quando text, usd real, tokens_in int, tokens_out int)")
        self.db.commit()

    def visto(self, caminho, tamanho, mtime):
        r = self.db.execute("select status from arq where caminho=? and tamanho=? and mtime=?", (caminho, tamanho, mtime)).fetchone()
        return r and r[0] in ("lido", "ignorado", "sem_texto")

    def marcar(self, caminho, tamanho, mtime, status, rid=None, info=""):
        self.db.execute("insert or replace into arq values(?,?,?,?,?,?)", (caminho, tamanho, mtime, status, rid, info[:500]))
        self.db.commit()

    def gasto_total(self):
        return self.db.execute("select coalesce(sum(usd),0) from gasto").fetchone()[0]

    def gastar(self, usd, ti, to):
        self.db.execute("insert into gasto(quando,usd,tokens_in,tokens_out) values(?,?,?,?)", (dt.datetime.now().isoformat(), usd, ti, to))
        self.db.commit()


# ---------------------------------------------------------------- Supabase
class Base:
    def __init__(self, email=None, seco=False):
        import requests
        self.rq = requests
        self.seco = seco
        self.sess = None
        if seco:
            return
        arq = HOME / "sessao.json"
        if arq.exists():
            try:
                s = json.loads(arq.read_text())
                self.sess = self._auth("token?grant_type=refresh_token", {"refresh_token": s["refresh_token"]})
            except Exception:
                self.sess = None
        if not self.sess:
            email = email or input("E-mail do Prognos: ").strip()
            senha = getpass.getpass("Senha do Prognos: ")
            self.sess = self._auth("token?grant_type=password", {"email": email, "password": senha})
        self._salvar()
        ok = self.post("rpc/sou_membro", {})
        if ok is not True:
            sys.exit("Este usuário não está autorizado no Prognos.")
        log(f"Conectado como {self.sess['user']['email']}")

    def _auth(self, path, body):
        r = self.rq.post(f"{SB_URL}/auth/v1/{path}", json=body, headers={"apikey": SB_KEY}, timeout=60)
        if r.status_code != 200:
            raise SystemExit(f"Falha no login: {r.text[:300]}")
        j = r.json()
        j["_exp"] = time.time() + j.get("expires_in", 3600) - 120
        return j

    def _salvar(self):
        (HOME / "sessao.json").write_text(json.dumps({"refresh_token": self.sess["refresh_token"]}))

    def _h(self, extra=None):
        if time.time() > self.sess["_exp"]:
            self.sess = self._auth("token?grant_type=refresh_token", {"refresh_token": self.sess["refresh_token"]})
            self._salvar()
        h = {"apikey": SB_KEY, "Authorization": "Bearer " + self.sess["access_token"], "Content-Type": "application/json"}
        h.update(extra or {})
        return h

    def _req(self, metodo, path, body=None, extra=None, tentativas=4):
        for t in range(tentativas):
            try:
                r = self.rq.request(metodo, f"{SB_URL}/rest/v1/{path}", data=None if body is None else json.dumps(body, default=str), headers=self._h(extra), timeout=120)
                if r.status_code in (429, 500, 502, 503, 504):
                    raise RuntimeError(f"{r.status_code} {r.text[:200]}")
                if r.status_code >= 400:
                    raise ValueError(f"{r.status_code} {r.text[:400]}")
                return r.json() if r.text else None
            except ValueError:
                raise
            except Exception as e:
                if t == tentativas - 1:
                    raise
                time.sleep(2 ** t)

    def post(self, path, body, extra=None):
        if self.seco:
            return None
        return self._req("POST", path, body, extra)

    def upsert_arquivo(self, row):
        if self.seco:
            return 0
        r = self._req("POST", "arquivos?on_conflict=unidade,caminho", [row], {"Prefer": "resolution=merge-duplicates,return=representation"})
        return r[0]["id"]

    def trocar_trechos(self, rid, trechos):
        if self.seco or not rid:
            return
        self._req("DELETE", f"trechos?arquivo_id=eq.{rid}")
        for i in range(0, len(trechos), 200):
            self._req("POST", "trechos", trechos[i:i + 200], {"Prefer": "return=minimal"})

    def eventos(self, evs):
        if self.seco or not evs:
            return
        for i in range(0, len(evs), 300):
            self._req("POST", "eventos?on_conflict=unidade,chave", evs[i:i + 300], {"Prefer": "resolution=ignore-duplicates,return=minimal"})

    def ia(self, prompt, tier="quick", max_tokens=4000, imagens=None):
        conteudo = []
        for b64 in imagens or []:
            conteudo.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}})
        conteudo.append({"type": "text", "text": prompt})
        for t in range(4):
            r = self.rq.post(f"{SB_URL}/functions/v1/ia", headers=self._h(), json={"tier": tier, "stream": False, "max_tokens": max_tokens, "messages": [{"role": "user", "content": conteudo}]}, timeout=300)
            if r.status_code in (429, 502, 503, 529):
                time.sleep(5 * (t + 1))
                continue
            if r.status_code != 200:
                raise RuntimeError(f"IA: {r.status_code} {r.text[:300]}")
            j = r.json()
            texto = "".join(b.get("text", "") for b in j.get("content", []) if b.get("type") == "text")
            return texto, j.get("usage", {})
        raise RuntimeError("IA indisponível depois de várias tentativas")


# ---------------------------------------------------------------- extração de texto
def texto_binario(dados: bytes) -> str:
    """Recupera texto legível de formatos antigos (.doc, .ppt) sem bibliotecas."""
    melhor = ""
    for cod in ("utf-16le", "latin-1"):
        t = dados.decode(cod, errors="ignore")
        runs = re.findall(r"[A-Za-z0-9À-ÖØ-öø-ý \t\r\n.,;:()/%°ºª\-+=]{25,}", t)
        s = "\n".join(r.strip() for r in runs if sum(c.isalpha() for c in r) > len(r) * .55 and len(set(r.strip())) > 8
                        and r.count(" ") >= len(r) / 14 and len(re.findall(r"[aeiouáéíóúãõâêô]", r.lower())) > len(r) * .22
                        and not re.search(r"(?:[A-Za-z0-9]{1,2}){0}(?:ABCDEF|CDEFGH|cdefgh|456789|stuvwx)", r))
        if len(s) > len(melhor):
            melhor = s
    return melhor


def ler_pdf(dados):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(stream=dados, filetype="pdf")
    paginas = [(i + 1, p.get_text("text")) for i, p in enumerate(doc)]
    return paginas, doc.page_count


def ler_tabela(nome, dados):
    ext = Path(nome).suffix.lower()
    linhas = []
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        for ws in wb.worksheets[:10]:
            for r in ws.iter_rows(values_only=True):
                if any(c not in (None, "") for c in r):
                    linhas.append([c for c in r])
                if len(linhas) > 200000:
                    break
    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=dados)
        for sh in wb.sheets()[:10]:
            for i in range(sh.nrows):
                row = []
                for c in sh.row(i):
                    if c.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate.xldate_as_datetime(c.value, wb.datemode))
                        except Exception:
                            row.append(c.value)
                    else:
                        row.append(c.value)
                linhas.append(row)
    else:
        import csv
        t = decodificar(dados)
        amostra = t[:5000]
        sep = max([";", ",", "\t", "|"], key=amostra.count)
        linhas = list(csv.reader(io.StringIO(t), delimiter=sep))
    return linhas


def decodificar(dados: bytes) -> str:
    if dados[:2] in (b"\xff\xfe", b"\xfe\xff") or b"\x00" in dados[:400]:
        return dados.decode("utf-16", errors="ignore")
    try:
        return dados.decode("utf-8")
    except UnicodeDecodeError:
        return dados.decode("latin-1", errors="ignore")


def cel(c):
    if isinstance(c, (dt.datetime, dt.date)):
        return c.strftime("%Y-%m-%d")
    return "" if c is None else str(c).strip()


def extrair(nome, dados):
    """Devolve (segmentos[(pagina, texto)], info, linhas_tabela|None)."""
    ext = Path(nome).suffix.lower()
    if ext == ".pdf":
        pags, n = ler_pdf(dados)
        return pags, f"{n} páginas", None
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(dados))
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                partes.append(" | ".join(c.text for c in r.cells if c.text.strip()))
        texto = "\n".join(x for x in partes if x.strip())
        if len(texto) < 200:  # texto em caixas de texto, cabeçalhos ou formas
            z = zipfile.ZipFile(io.BytesIO(dados))
            xmls = [n for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
            texto = "\n".join(" ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", z.read(n).decode("utf-8", "ignore"))) for n in xmls)
        return [(None, texto)], "", None
    if ext == ".pptx":
        from pptx import Presentation
        pr = Presentation(io.BytesIO(dados))
        segs = []
        for i, s in enumerate(pr.slides, 1):
            tx = [sh.text_frame.text for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame]
            segs.append((i, "\n".join(tx)))
        return segs, f"{len(segs)} slides", None
    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
        linhas = ler_tabela(nome, dados)
        texto = "\n".join(" | ".join(cel(c) for c in r if cel(c)) for r in linhas[:3000])
        return [(None, texto)], f"{len(linhas)} linhas", linhas
    if ext == ".msg":
        import extract_msg
        m = extract_msg.Message(io.BytesIO(dados))
        return [(None, f"Assunto: {m.subject}\nDe: {m.sender}\nData: {m.date}\n\n{m.body or ''}")], "e-mail", None
    if ext == ".eml":
        import email
        from email import policy
        m = email.message_from_bytes(dados, policy=policy.default)
        corpo = m.get_body(preferencelist=("plain", "html"))
        t = corpo.get_content() if corpo else ""
        return [(None, f"Assunto: {m['subject']}\nDe: {m['from']}\nData: {m['date']}\n\n{re.sub('<[^>]+>', ' ', t)}")], "e-mail", None
    if ext in (".doc", ".ppt", ".rtf", ".wri", ".pps"):
        return [(None, texto_binario(dados))], "formato antigo (texto recuperado)", None
    if ext in TEXTO:
        t = decodificar(dados)
        if ext in (".htm", ".html", ".xml") and ext != ".l5x":
            t = re.sub(r"<(script|style)[\s\S]*?</\1>", " ", t, flags=re.I)
            t = re.sub(r"<[^>]+>", " ", t)
        return [(None, t)], "", None
    if ext == ".dxf":
        try:
            import ezdxf
            d = ezdxf.read(io.StringIO(decodificar(dados)))
            tx = [e.dxf.text for e in d.modelspace().query("TEXT")] + [e.text for e in d.modelspace().query("MTEXT")]
            return [(None, "\n".join(tx))], "textos do desenho", None
        except Exception:
            return [(None, "")], "dxf sem texto legível", None
    return [], "", None


def resumo_desenho(caminho: str):
    nome = Path(caminho).name
    m = re.search(r"(DE|DES|FD|LD)[-_ ]?([A-Z]?\d{6,10})[-_ ]?(\d{4})?([A-Z]{3})(\d{3,4})", nome, re.I)
    partes = [f"Desenho técnico {nome}", f"Pasta: {Path(caminho).parent}"]
    doc = None
    if m:
        doc = m.group(0).upper()
        disc = DISCIPLINA.get(m.group(4).upper(), m.group(4).upper())
        partes.append(f"Número {doc}; projeto {m.group(2)}; área {m.group(3) or '-'}; disciplina {disc}")
    return " · ".join(partes), doc


def trechos(segs, maximo):
    out = []
    for pag, t in segs:
        t = re.sub(r"[ \t]+", " ", (t or "")).strip()
        t = re.sub(r"\n{3,}", "\n\n", t)
        i = 0
        while i < len(t) and len(out) < maximo:
            pedaco = t[i:i + 1500]
            if i + 1500 < len(t):
                corte = max(pedaco.rfind("\n"), pedaco.rfind(". "))
                if corte > 900:
                    pedaco = pedaco[:corte + 1]
            if len(pedaco.strip()) > 40:
                out.append((pag, pedaco.strip()))
            i += max(200, len(pedaco) - 150)
    return out


# ---------------------------------------------------------------- tabelas SAP e históricos
TIPO_RE = [("troca_preventiva", r"preventiv"), ("troca_corretiva", r"troc|substitu"), ("falha", r"quebr|falh|travad|romp|queim|desarm|curto|parad[ao] n|emerg"),
           ("inspecao", r"inspe|rota|verifica"), ("medicao", r"medi[cç]|mm/s|termograf|an[aá]lise de [oó]leo"), ("reparo", r"repar|solda|recupera|ajust")]


def data_de(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", s)
    if m:
        a = int(m.group(3)); a += 2000 if a < 100 else 0
        if int(m.group(2)) <= 12:
            return f"{a}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return None


def eventos_de_tabela(linhas, caminho):
    """Exportações de notas e ordens do SAP (IW29, IW39, IW47 etc.) viram eventos sem custo de IA."""
    if not linhas:
        return []
    hi = 0
    for i, r in enumerate(linhas[:20]):
        if sum(1 for c in r if isinstance(c, str) and c.strip() and not c.strip().replace(",", ".").replace(".", "").isdigit()) >= 3:
            hi = i
            break
    cab = [cel(c).lower() for c in linhas[hi]]
    def col(*pad):
        for p in pad:
            for j, h in enumerate(cab):
                if re.search(p, h):
                    return j
        return -1
    cd = col(r"data|date|in[ií]cio|criad|fim|conclus")
    cn, co = col(r"^nota|n[ºo°]?\s*nota|notifica"), col(r"^ordem|n[ºo°]?\s*ordem|\bom\b|order")
    ce = col(r"equipament|equip|local de inst|loc\.? ?inst|tag|ativo|func.*loc")
    ctx = [j for j, h in enumerate(cab) if re.search(r"descri|texto|text|observ|sintoma|problema|causa|a[cç][aã]o|servi[cç]o|relato|breve", h)]
    ctp = col(r"tipo de nota|tipo de ordem|^tipo|tp\.? ?nota|tp\.? ?ordem|categoria")
    if cd < 0 or (cn < 0 and co < 0 and ce < 0) or not ctx:
        return []
    evs = []
    for r in linhas[hi + 1:]:
        if len(r) <= max(cd, 0):
            continue
        d = data_de(r[cd])
        if not d:
            continue
        txt = " / ".join(cel(r[j]) for j in ctx if j < len(r) and cel(r[j]))[:500]
        if not txt:
            continue
        tp = (cel(r[ctp]) + " " if 0 <= ctp < len(r) else "") + txt
        tipo = next((t for t, p in TIPO_RE if re.search(p, tp, re.I)), "outro")
        nota = cel(r[cn]) if 0 <= cn < len(r) else ""
        ordem = cel(r[co]) if 0 <= co < len(r) else ""
        equip = cel(r[ce]) if 0 <= ce < len(r) else ""
        evs.append({"data": d, "tipo_evento": tipo, "equipamento": equip[:200] or None, "texto": txt, "fonte": f"SAP: {Path(caminho).name}",
                    "nota_sap": nota or None, "ordem_sap": ordem or None, "chave": f"sap|{nota or ordem}|{d}|{txt[:40]}".lower()})
    return evs


PROMPT_EVENTOS = """Você é o extrator de eventos de manutenção do Prognos. Do texto abaixo (arquivo "{nome}"), extraia somente eventos reais e datados de equipamentos: falhas, trocas corretivas, trocas preventivas, reparos, inspeções, medições e sintomas. Não invente datas; se a data não estiver no texto, ignore o evento. Identifique o equipamento ou posição do jeito que aparece no texto.
Responda só JSON: {{"eventos":[{{"data":"AAAA-MM-DD","tipo_evento":"falha|troca_corretiva|troca_preventiva|reparo|inspecao|medicao|sintoma|outro","equipamento":"","componente":"","modo":"","texto":"até 200 caracteres"}}]}}
Se não houver eventos, responda {{"eventos":[]}}.

TEXTO:
{texto}"""


def json_de(t):
    try:
        return json.loads(t)
    except Exception:
        pass
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", t)
    if m:
        try:
            return json.loads(m.group(1))
        except Exception:
            pass
    a, b = t.find("{"), t.rfind("}")
    if a >= 0 and b > a:
        try:
            return json.loads(t[a:b + 1])
        except Exception:
            pass
    return None


# ---------------------------------------------------------------- principal
def categoria(rel: Path):
    partes = rel.parts[:-1]
    return " / ".join(partes[:2]) if partes else "(raiz)"


def iterar(pasta: Path, limite):
    n = 0
    for raiz, dirs, arqs in os.walk(pasta):
        dirs[:] = [d for d in dirs if not d.startswith((".", "$")) and d.lower() not in ("system volume information", "recycler", "$recycle.bin")]
        for a in arqs:
            p = Path(raiz) / a
            try:
                st = p.stat()
            except OSError:
                continue
            yield p, st.st_size, st.st_mtime
            n += 1
            if limite and n >= limite:
                return


def membros_zip(dados, nome):
    """Arquivos legíveis dentro de um ZIP (um nível)."""
    try:
        z = zipfile.ZipFile(io.BytesIO(dados))
    except Exception:
        return []
    out = []
    for info in z.infolist():
        if info.is_dir() or info.file_size > 150 * 1024 * 1024:
            continue
        ext = Path(info.filename).suffix.lower()
        if ext in IGNORAR or ext in IMAGEM or ext in VIDEO or ext in COMPACTADO:
            continue
        try:
            out.append((f"{nome}!/{info.filename}", z.read(info)))
        except Exception:
            pass
        if len(out) >= 300:
            break
    return out


def principal():
    ap = argparse.ArgumentParser(description="Ingestão do arquivo técnico no Prognos")
    ap.add_argument("pasta", nargs="?", default=r"G:\1 - Arquivo Tecnico")
    ap.add_argument("--email")
    ap.add_argument("--inventario", action="store_true", help="só conta arquivos, tamanhos e tipos; não envia nada")
    ap.add_argument("--seco", action="store_true", help="lê e extrai, mas não envia nada para a base")
    ap.add_argument("--limite", type=int, default=0, help="processa no máximo N arquivos")
    ap.add_argument("--max-mb", type=int, default=150, help="ignora arquivos maiores que isto (MB), exceto ZIP")
    ap.add_argument("--max-trechos", type=int, default=250, help="máximo de trechos indexados por arquivo")
    ap.add_argument("--ia", action="store_true", help="usa a IA para extrair eventos de relatórios e históricos")
    ap.add_argument("--orcamento-usd", type=float, default=5.0, help="teto de gasto com IA nesta pasta (estimado)")
    ap.add_argument("--preco-entrada", type=float, default=1.0, help="US$ por milhão de tokens de entrada do modelo rápido (confira na Anthropic)")
    ap.add_argument("--preco-saida", type=float, default=5.0, help="US$ por milhão de tokens de saída do modelo rápido (confira na Anthropic)")
    ap.add_argument("--refazer", action="store_true", help="relê arquivos já processados")
    a = ap.parse_args()

    pasta = Path(a.pasta)
    if not pasta.exists():
        sys.exit(f"Pasta não encontrada: {pasta}")

    if a.inventario:
        tipos, tot, n = {}, 0, 0
        for p, tam, _ in iterar(pasta, a.limite):
            e = p.suffix.lower() or "(sem)"
            c = tipos.setdefault(e, [0, 0]); c[0] += 1; c[1] += tam; tot += tam; n += 1
            if n % 5000 == 0:
                log(f"  {n} arquivos…")
        log(f"\n{n} arquivos · {tot / 1e9:.1f} GB em {pasta}\n")
        log(f"{'tipo':10} {'arquivos':>9} {'GB':>9}  tratamento")
        for e, (q, t) in sorted(tipos.items(), key=lambda x: -x[1][1])[:40]:
            trat = "ignorado" if e in IGNORAR else "imagem: só nome" if e in IMAGEM else "desenho: nome e número" if e in DESENHO else "vídeo: só nome" if e in VIDEO else "compactado" if e in COMPACTADO else "texto indexado"
            log(f"{e:10} {q:9d} {t / 1e9:9.2f}  {trat}")
        return

    est = Estado(pasta)
    base = Base(a.email, a.seco)
    lim_bytes = a.max_mb * 1024 * 1024
    stats = dict(lidos=0, pulados=0, ignorados=0, erros=0, trechos=0, eventos=0, bytes=0)
    gasto = est.gasto_total()
    inicio = time.time()

    def processar(caminho_rel, nome, dados, tamanho, mtime, rel_path):
        nonlocal gasto
        ext = Path(nome).suffix.lower()
        row = {"caminho": caminho_rel, "nome": Path(nome).name, "ext": ext, "pasta": str(Path(caminho_rel).parent), "tamanho": tamanho,
               "modificado": dt.datetime.fromtimestamp(mtime).isoformat(), "categoria": categoria(rel_path), "status": "lido", "processado_em": dt.datetime.now().isoformat()}
        evs, segs, info, linhas = [], [], "", None
        if ext in DESENHO or ext in IMAGEM or ext in VIDEO:
            txt, doc = resumo_desenho(caminho_rel) if ext in DESENHO else (f"{'Imagem' if ext in IMAGEM else 'Vídeo'} {Path(nome).name} · Pasta: {Path(caminho_rel).parent}", None)
            segs = [(None, txt)]
            row["documento"] = doc
            if ext == ".dxf" and dados:
                s2, info, _ = extrair(nome, dados)
                segs += s2
        else:
            segs, info, linhas = extrair(nome, dados)
        ts = trechos(segs, a.max_trechos)
        row["caracteres"] = sum(len(t) for _, t in ts)
        row["trechos"] = len(ts)
        row["paginas"] = sum(1 for p, _ in segs if p) or None
        if not ts:
            row["status"] = "sem_texto"
            row["erro"] = info or "sem texto legível (digitalizado? use OCR depois)"
        rid = base.upsert_arquivo(row)
        base.trocar_trechos(rid, [{"arquivo_id": rid, "ordem": i, "pagina": p, "texto": t} for i, (p, t) in enumerate(ts)])
        if linhas:
            evs = eventos_de_tabela(linhas, caminho_rel)
        texto_total = "\n".join(t for _, t in segs)
        if a.ia and not evs and texto_total and HIST_RE.search(texto_total) and not MANUAL_PASTAS.search(caminho_rel):
            for i in range(0, min(len(texto_total), 60000), 20000):
                bloco = texto_total[i:i + 20000]
                est_in = len(bloco) / 3.5 + 600
                custo_prev = est_in / 1e6 * a.preco_entrada + 1500 / 1e6 * a.preco_saida
                if gasto + custo_prev > a.orcamento_usd:
                    log(f"  teto de IA atingido (US$ {gasto:.2f}); continuo só indexando texto")
                    a.ia = False
                    break
                try:
                    resp, uso = base.ia(PROMPT_EVENTOS.format(nome=Path(nome).name, texto=bloco))
                    ti, to = uso.get("input_tokens", est_in), uso.get("output_tokens", 800)
                    c = ti / 1e6 * a.preco_entrada + to / 1e6 * a.preco_saida
                    gasto += c
                    est.gastar(c, ti, to)
                    for e in (json_de(resp) or {}).get("eventos", []):
                        d = data_de(e.get("data"))
                        if d:
                            evs.append({"data": d, "tipo_evento": e.get("tipo_evento") or "outro", "equipamento": (e.get("equipamento") or "")[:200] or None, "componente": e.get("componente") or None,
                                        "modo": e.get("modo") or None, "texto": (e.get("texto") or "")[:500], "fonte": Path(nome).name, "chave": f"ia|{caminho_rel}|{d}|{(e.get('texto') or '')[:40]}".lower()})
                except Exception as ex:
                    log(f"  IA falhou em {nome}: {ex}")
                    break
        for e in evs:
            e["arquivo_id"] = rid or None
        base.eventos(evs)
        stats["trechos"] += len(ts)
        stats["eventos"] += len(evs)
        return rid, row["status"], info

    try:
        for p, tam, mt in iterar(pasta, a.limite):
            rel = p.relative_to(pasta)
            caminho = str(p)
            ext = p.suffix.lower()
            if not a.refazer and est.visto(caminho, tam, mt):
                stats["pulados"] += 1
                continue
            if ext in IGNORAR or (tam > lim_bytes and ext not in COMPACTADO and ext not in DESENHO and ext not in IMAGEM and ext not in VIDEO):
                est.marcar(caminho, tam, mt, "ignorado", info="tipo ignorado ou grande demais")
                stats["ignorados"] += 1
                continue
            try:
                if ext in DESENHO and ext != ".dxf" or ext in IMAGEM or ext in VIDEO:
                    rid, status, info = processar(caminho, p.name, None, tam, mt, rel)
                elif ext == ".zip":
                    dados = p.read_bytes() if tam <= 2 * 1024 ** 3 else b""
                    rid, status, info = processar(caminho, p.name, None, tam, mt, rel)
                    for nome_m, dm in membros_zip(dados, caminho):
                        processar(nome_m, nome_m.split("!/")[-1], dm, len(dm), mt, rel)
                elif ext in (".7z", ".rar"):
                    rid, status, info = processar(caminho, p.name, None, tam, mt, rel)
                    info = "compactado: extraia para indexar o conteúdo"
                else:
                    dados = p.read_bytes()
                    rid, status, info = processar(caminho, p.name, dados, tam, mt, rel)
                est.marcar(caminho, tam, mt, status, rid, info or "")
                stats["lidos"] += 1
                stats["bytes"] += tam
            except KeyboardInterrupt:
                raise
            except Exception as ex:
                stats["erros"] += 1
                est.marcar(caminho, tam, mt, "erro", info=str(ex))
                log(f"  erro em {rel}: {str(ex)[:160]}")
            feitos = stats["lidos"] + stats["erros"]
            if feitos and feitos % 25 == 0:
                vel = stats["bytes"] / 1e6 / max(1, time.time() - inicio)
                log(f"{feitos} arquivos · {stats['trechos']} trechos · {stats['eventos']} eventos · {vel:.1f} MB/s · IA US$ {gasto:.2f} · {str(rel)[:70]}")
    except KeyboardInterrupt:
        log("\nInterrompido. Rode de novo para continuar de onde parou.")
    log(f"\nConcluído: {stats['lidos']} lidos, {stats['pulados']} já estavam na base, {stats['ignorados']} ignorados, {stats['erros']} com erro.")
    log(f"Trechos indexados: {stats['trechos']} · eventos: {stats['eventos']} · gasto estimado com IA: US$ {gasto:.2f}")
    if a.seco:
        log("Modo seco: nada foi enviado.")


if __name__ == "__main__":
    principal()
